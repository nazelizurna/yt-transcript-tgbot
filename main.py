"""
YouTube Transcript -> Telegram Document Bot (Webhook Architecture)
====================================================================
"""

import os
import re
import tempfile
import logging
import traceback
from typing import Optional

import requests
from fastapi import FastAPI, Request, HTTPException
from docx import Document
from docx.shared import Pt

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("yt-transcript-bot")

app = FastAPI()

TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN")
WEBHOOK_SECRET = os.environ.get("WEBHOOK_SECRET", "")
TELEGRAM_API_BASE = "https://api.telegram.org/bot{token}"

YOUTUBE_ID_PATTERNS = [
    r"(?:youtube\.com/watch\?v=)([A-Za-z0-9_-]{11})",
    r"(?:youtube\.com/shorts/)([A-Za-z0-9_-]{11})",
    r"(?:youtube\.com/embed/)([A-Za-z0-9_-]{11})",
    r"(?:youtu\.be/)([A-Za-z0-9_-]{11})",
    r"(?:m\.youtube\.com/watch\?v=)([A-Za-z0-9_-]{11})",
]


def extract_video_id(text: str) -> Optional[str]:
    if not text:
        return None
    for pattern in YOUTUBE_ID_PATTERNS:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return None


def tg_call(method: str, **kwargs):
    if not TELEGRAM_BOT_TOKEN:
        raise RuntimeError("TELEGRAM_BOT_TOKEN not configured")
    url = f"{TELEGRAM_API_BASE.format(token=TELEGRAM_BOT_TOKEN)}/{method}"
    files = kwargs.pop("files", None)
    resp = requests.post(url, data=kwargs, files=files, timeout=30)
    if not resp.ok:
        logger.error(f"Telegram API {method} failed: {resp.status_code} {resp.text}")
    resp.raise_for_status()
    return resp.json()


def send_message(chat_id: int, text: str) -> Optional[int]:
    try:
        result = tg_call("sendMessage", chat_id=chat_id, text=text)
        return result["result"]["message_id"]
    except Exception:
        logger.error(f"send_message failed:\n{traceback.format_exc()}")
        return None


def delete_message(chat_id: int, message_id: int) -> None:
    try:
        tg_call("deleteMessage", chat_id=chat_id, message_id=message_id)
    except Exception as e:
        logger.warning(f"Could not delete status message: {e}")


def send_document(chat_id: int, file_path: str, caption: str) -> None:
    with open(file_path, "rb") as f:
        tg_call("sendDocument", chat_id=chat_id, caption=caption, files={"document": f})


def _parse_vtt(vtt_text: str) -> list:
    """Parse a WEBVTT subtitle file into [{start, text}, ...] segments."""
    segments = []
    time_pattern = re.compile(r"(\d{2}:\d{2}:\d{2}\.\d{3}) --> ")
    tag_pattern = re.compile(r"<[^>]+>")

    current_start = None
    current_text_lines = []

    def flush():
        if current_start is not None and current_text_lines:
            text = " ".join(current_text_lines).strip()
            text = tag_pattern.sub("", text)
            if text:
                segments.append({"start": current_start, "text": text})

    for line in vtt_text.splitlines():
        match = time_pattern.match(line)
        if match:
            flush()
            h, m, s = match.group(1).split(":")
            current_start = int(h) * 3600 + int(m) * 60 + float(s)
            current_text_lines = []
        elif (
            line.strip()
            and not line.strip().isdigit()
            and "WEBVTT" not in line
            and "-->" not in line
        ):
            current_text_lines.append(line.strip())
    flush()
    return segments


def _parse_json3(data: dict) -> list:
    """Parse YouTube's json3 caption format into [{start, text}, ...] segments."""
    segments = []
    for event in data.get("events", []):
        if "segs" not in event:
            continue
        start = event.get("tStartMs", 0) / 1000.0
        text = "".join(seg.get("utf8", "") for seg in event["segs"]).strip()
        if text:
            segments.append({"start": start, "text": text})
    return segments


RAPIDAPI_KEY = os.environ.get("RAPIDAPI_KEY")
RAPIDAPI_HOST = "youtube-transcripts.p.rapidapi.com"  # confirm exact value from your RapidAPI code snippet
RAPIDAPI_URL = f"https://{RAPIDAPI_HOST}/youtube/transcript"  # confirm exact path from your RapidAPI code snippet


def fetch_transcript_rapidapi(video_id: str) -> Optional[list]:
    """
    Primary transcript source: Supadata's YouTube Transcripts API via RapidAPI.
    Runs from Supadata's infrastructure, not Render's IP, so it sidesteps
    the datacenter-IP block entirely.
    """
    if not RAPIDAPI_KEY:
        return None

    try:
        resp = requests.get(
            RAPIDAPI_URL,
            headers={
                "X-RapidAPI-Key": RAPIDAPI_KEY,
                "X-RapidAPI-Host": RAPIDAPI_HOST,
            },
            params={"url": f"https://www.youtube.com/watch?v={video_id}"},
            timeout=30,
        )
        resp.raise_for_status()
        data = resp.json()

        # Supadata's response shape: {"lang": "en", "content": [{"text","offset","duration"}, ...]}
        raw_segments = data.get("content") or data.get("transcript") or []
        if not raw_segments:
            return None

        segments = [
            {"start": seg["offset"] / 1000.0, "text": seg["text"]}
            for seg in raw_segments
            if seg.get("text")
        ]
        return segments or None

    except Exception:
        logger.error(f"RapidAPI transcript fetch failed for {video_id}:\n{traceback.format_exc()}")
        return None


def fetch_transcript_ytdlp(video_id: str) -> Optional[list]:
    """Fallback: original yt-dlp method (kept as-is, renamed)."""
    try:
        import yt_dlp

        url = f"https://www.youtube.com/watch?v={video_id}"
        ydl_opts = {
            "quiet": False,
            "no_warnings": False,
            "verbose": True,
            "skip_download": True,
            "writeautomaticsub": True,
            "writesubtitles": True,
        }

        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            info = ydl.extract_info(url, download=False)

        manual_subs = info.get("subtitles") or {}
        auto_subs = info.get("automatic_captions") or {}

        preferred_langs = [
            "en", "en-US", "en-GB", "ru", "es", "pt", "fr", "de",
            "hi", "id", "ar", "ja", "ko", "zh-Hans", "zh-Hant", "it", "tr", "vi",
        ]

        track_list = None
        for lang in preferred_langs:
            if lang in manual_subs:
                track_list = manual_subs[lang]
                break
        if track_list is None:
            for lang in preferred_langs:
                if lang in auto_subs:
                    track_list = auto_subs[lang]
                    break
        if track_list is None:
            combined = {**auto_subs, **manual_subs}
            if combined:
                track_list = next(iter(combined.values()))

        if not track_list:
            logger.error(f"No caption tracks found for {video_id}")
            return None

        fmt_entry = next((f for f in track_list if f.get("ext") == "json3"), None)
        if fmt_entry is None:
            fmt_entry = next(
                (f for f in track_list if f.get("ext") == "vtt"), track_list[0]
            )

        resp = requests.get(fmt_entry["url"], timeout=30)
        resp.raise_for_status()

        if fmt_entry.get("ext") == "json3":
            segments = _parse_json3(resp.json())
        else:
            segments = _parse_vtt(resp.text)

        return segments or None

    except Exception:
        logger.error(f"yt-dlp transcript fetch failed for {video_id}:\n{traceback.format_exc()}")
        return None


def fetch_transcript(video_id: str) -> Optional[list]:
    """Try RapidAPI (Supadata) first — off-Render infra, dodges the IP block.
    Fall back to yt-dlp only if RapidAPI fails or the free quota is exhausted."""
    segments = fetch_transcript_rapidapi(video_id)
    if segments:
        return segments

    logger.warning(f"RapidAPI transcript unavailable for {video_id}, falling back to yt-dlp")
    return fetch_transcript_ytdlp(video_id)


def format_timestamp(seconds: float) -> str:
    total = int(seconds)
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    if h > 0:
        return f"[{h:02d}:{m:02d}:{s:02d}]"
    return f"[{m:02d}:{s:02d}]"


def build_docx(video_id: str, segments: list) -> str:
    doc = Document()
    doc.add_heading(f"YouTube Transcript - Video {video_id}", level=1)

    for seg in segments:
        p = doc.add_paragraph()
        ts_run = p.add_run(format_timestamp(seg["start"]) + " ")
        ts_run.bold = True
        ts_run.font.size = Pt(11)
        text_run = p.add_run(seg["text"])
        text_run.font.size = Pt(11)

    fd, path = tempfile.mkstemp(suffix=".docx", dir="/tmp", prefix=f"{video_id}_")
    os.close(fd)
    doc.save(path)
    return path


def process_update(update: dict) -> None:
    message = update.get("message") or update.get("edited_message")
    if not message:
        return

    chat_id = message["chat"]["id"]
    text = message.get("text", "")

    video_id = extract_video_id(text)
    if not video_id:
        send_message(chat_id, "❌ Please send a valid YouTube link.")
        return

    status_id = send_message(chat_id, "⏳ Extracting video transcript...")

    file_path = None
    try:
        segments = fetch_transcript(video_id)
        if not segments:
            send_message(chat_id, "❌ Transcript unavailable for this video.")
            return

        file_path = build_docx(video_id, segments)
        send_document(chat_id, file_path, caption=f"📄 File ready for Video ID: {video_id}")
    except Exception:
        logger.error(f"process_update failed for video {video_id}:\n{traceback.format_exc()}")
        send_message(chat_id, "❌ Something went wrong generating your transcript.")
    finally:
        if status_id:
            delete_message(chat_id, status_id)
        if file_path and os.path.exists(file_path):
            os.remove(file_path)


@app.post("/webhook/{secret}")
async def telegram_webhook(secret: str, request: Request):
    # Never 500 on webhook secret mismatch checks or config issues —
    # always return 200 so Telegram doesn't hammer retries, and log everything.
    if WEBHOOK_SECRET and secret != WEBHOOK_SECRET:
        logger.warning(f"Rejected webhook call with bad secret: {secret}")
        return {"ok": True}

    if not TELEGRAM_BOT_TOKEN:
        logger.error("TELEGRAM_BOT_TOKEN is NOT SET in this environment — check Render Environment tab.")
        return {"ok": True}

    try:
        update = await request.json()
    except Exception:
        logger.error(f"Failed to parse incoming update as JSON:\n{traceback.format_exc()}")
        return {"ok": True}

    try:
        process_update(update)
    except Exception:
        logger.error(f"Unhandled error in process_update:\n{traceback.format_exc()}")

    return {"ok": True}


@app.get("/")
async def health_check():
    return {"status": "ok"}


@app.get("/debug")
async def debug():
    """Temporary: confirms whether env vars are actually loaded in this deployment."""
    return {
        "token_set": bool(TELEGRAM_BOT_TOKEN),
        "token_length": len(TELEGRAM_BOT_TOKEN) if TELEGRAM_BOT_TOKEN else 0,
        "webhook_secret_set": bool(WEBHOOK_SECRET),
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.environ.get("PORT", 8000)))
